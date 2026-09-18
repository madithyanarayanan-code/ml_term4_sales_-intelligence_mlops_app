from __future__ import annotations
import io
import os
import requests
import pandas as pd


def _deterministic_report(result):
    rfm = result["rfm"]
    cleaned = result["cleaned"]
    training = result["training"]
    assoc = result["association"]
    stages = result["stages"]
    recs = result["recommendations"]
    lines = [
        "# Sales Customer Analysis Report",
        "",
        "## Executive decision summary",
        f"- Transactions analysed: {len(cleaned):,}",
        f"- Customers analysed: {len(rfm):,}",
        f"- Total sales: {cleaned['AnalyticsAmount'].sum():,.2f}",
        f"- Estimated churn: {rfm['Churn'].mean():.1%}",
        f"- Customer reference period: {stages['5'].get('reference_date')}",
        "",
        "## Customer intelligence",
        f"Average customer monetary value is {rfm['Monetary'].mean():,.2f}; median recency is {rfm['Recency'].median():.0f} days and average purchase frequency is {rfm['Frequency'].mean():.2f} orders.",
        "",
        "## Model results",
    ]
    if training.get("classification"):
        best = max(training["classification"], key=lambda x:x["f1"])
        lines.append(f"Best churn classifier by F1: **{best['model']}** (F1 {best['f1']:.3f}, ROC-AUC {best['roc_auc']:.3f} where available).")
    else:
        lines.append("Churn classification could not be evaluated because the data did not contain enough class variation/customer observations.")
    if training.get("regression"):
        best = max(training["regression"], key=lambda x:x["r2"])
        lines.append(f"Best monetary-value regressor by R²: **{best['model']}** (R² {best['r2']:.3f}, adjusted R² {best['adjusted_r2']:.3f}, RMSE {best['rmse']:,.2f}, MAE {best['mae']:,.2f}).")
    else:
        lines.append("Regression evaluation was skipped because the customer-level sample was insufficient.")
    lines += ["", "## Association insights"]
    if assoc.get("available") and not assoc.get("rules", pd.DataFrame()).empty:
        top = assoc["rules"].iloc[0]
        lines.append(f"The strongest observed association has lift {top['lift']:.2f} and confidence {top['confidence']:.1%}.")
    else:
        lines.append("No sufficiently strong association rules were produced under the current thresholds.")
    lines += ["", "## Recommendations"] + [f"- {x}" for x in recs]
    lines += ["", "## Governance and reproducibility", "The workflow follows the eight defined lifecycle stages. Raw/processed/engineered/evaluation artifacts are logically separated, model candidates are compared before selection, and developer diagnostics are isolated from the main business views."]
    return "\n".join(lines)


def _llama_polish(draft: str, model: str) -> str:
    """Optional local Llama integration via Ollama. Falls back to draft if unavailable."""
    url = os.getenv("OLLAMA_URL", "http://localhost:11434/api/generate")
    prompt = f"""You are a management analyst. Rewrite the following sales analytics report into a concise, evidence-based executive report. Preserve all numeric values. Do not invent facts. Use headings, bullets and clear business actions.\n\nREPORT:\n{draft}"""
    try:
        r = requests.post(url, json={"model": model, "prompt": prompt, "stream": False}, timeout=90)
        r.raise_for_status()
        text = r.json().get("response", "").strip()
        return text or draft
    except Exception:
        return draft


def generate_report(result, model="llama3.1", as_docx=False):
    draft = _deterministic_report(result)
    content = _llama_polish(draft, model)
    if not as_docx:
        return content
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    for line in content.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.size = Pt(10.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
